import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Modern Dark Cyber Threat Intelligence
    C_BG = RGBColor(11, 15, 25)         # Deep Midnight Navy
    C_CARD = RGBColor(22, 29, 47)       # Slate Card
    C_CARD_HOVER = RGBColor(30, 41, 69) # Elevated Card
    C_CYAN = RGBColor(56, 189, 248)     # Accent Primary
    C_RED = RGBColor(248, 113, 113)     # Attack / Exploit / Danger
    C_GREEN = RGBColor(74, 222, 128)    # Defense / Secure / Patch
    C_YELLOW = RGBColor(250, 204, 21)   # Warning / Reflected
    C_PURPLE = RGBColor(192, 132, 252)  # Data Leak / Secret Flag
    C_TEXT = RGBColor(241, 245, 249)    # Bright Text
    C_MUTED = RGBColor(148, 163, 184)   # Slate Subtitle

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        # Tag Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.2), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(30, 41, 59)
        badge.line.color.rgb = C_CYAN
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.margin_top = Inches(0.04)
        p_b = tf_b.paragraphs[0]
        p_b.text = tag_text
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = C_CYAN
        p_b.alignment = PP_ALIGN.CENTER

        # Main Title & Subtitle
        tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(1.0))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(13)
            p2.font.color.rgb = C_MUTED
            p2.space_before = Pt(4)

    # -----------------------------------------------------------------
    # SLIDE 1: Title & Hero
    # -----------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Left accent glow bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.8), Inches(0.12), Inches(4.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_CYAN
    bar.line.fill.background()

    # Main Hero Box
    h_box = s1.shapes.add_textbox(Inches(1.4), Inches(1.6), Inches(10.8), Inches(4.5))
    tf1 = h_box.text_frame
    tf1.word_wrap = True

    p_tag = tf1.paragraphs[0]
    p_tag.text = "[ HANDS-ON CYBERSECURITY & THREAT SIMULATION ]"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = C_CYAN

    p_title = tf1.add_paragraph()
    p_title.text = "SecLab: Real-World Web Exploit & Defense"
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = C_TEXT
    p_title.space_before = Pt(8)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Simulasi Kebocoran Data Perusahaan (Data Breach), SQLi Auth Bypass, Triad XSS, Web Shell RCE & Remediasi Kode Aman"
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = C_YELLOW
    p_desc.space_before = Pt(12)

    p_env = tf1.add_paragraph()
    p_env.text = "Engine: SecLab Enterprise | Database: MySQL (seclab_db) via phpMyAdmin | Mode: Vulnerable vs Secure"
    p_env.font.size = Pt(13)
    p_env.font.color.rgb = C_MUTED
    p_env.space_before = Pt(28)

    # -----------------------------------------------------------------
    # SLIDE 2: Arsitektur & Skenario Data Breach Real
    # -----------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "MODUL ARCHITECTURE", "Topologi Sistem Lab & Data Sensitif MySQL", "Mengapa simulasi ini merefleksikan insiden kebocoran data di dunia nyata")

    cards_s2 = [
        ("Database MySQL seclab_db", "Tersambung langsung ke phpMyAdmin lokal. Berisi tabel users (kredensial nyata, NIK KTP, Kartu Kredit, Gaji, dan API Token).", C_PURPLE, "Real Database Storage"),
        ("Login Portal (Auth Bypass)", "Autentikasi login karyawan dengan celah SQL Injection. Penyerang dapat masuk tanpa password sebagai Super Administrator.", C_RED, "Entry Gate Attack"),
        ("Direktori Profil & UNION SQLi", "Endpoint pencarian karyawan internal yang rentan terhadap exfiltration seluruh data rahasia perusahaan.", C_YELLOW, "Data Exfiltration"),
        ("Live DB Inspector & Audit", "Fitur verifikasi langsung untuk membandingkan data hasil eksploitasi dengan rekaman fisik di phpMyAdmin.", C_GREEN, "Proof of Leak Audit")
    ]

    for i, (title, desc, color, tag) in enumerate(cards_s2):
        col = i % 2
        row = i // 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.9 + row * 2.5)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = f"[{tag}] {title}"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT
        p2.space_before = Pt(8)

    # -----------------------------------------------------------------
    # SLIDE 3: SQLi Part 1 - Login Authentication Bypass
    # -----------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "ATTACK VECTOR 1", "SQL Injection: Login Authentication Bypass (login.php)", "Bagaimana penyerang melewati proteksi password tanpa kredensial yang sah")

    # Left Box: Exploit Anatomy
    c_l = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.7), Inches(5.0))
    c_l.fill.solid()
    c_l.fill.fore_color.rgb = C_CARD
    c_l.line.color.rgb = C_RED
    c_l.line.width = Pt(1.5)

    tf_l = c_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = Inches(0.25)

    p = tf_l.paragraphs[0]
    p.text = "Anatomi Kerentanan pada login.php"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_RED

    points_l = [
        "Kueri Rentan:\nSELECT * FROM users WHERE username = '$u' AND password = '$p'",
        "Input Penyerang (Username):\nsuperadmin' #",
        "Kueri yang Dijalankan MySQL:\nSELECT * FROM users WHERE username = 'superadmin' #' AND password = '...'",
        "Hasil: Tanda '#' memotong verifikasi password. Database mengembalikan user admin, sesi dibuat, login sukses 100%!"
    ]
    for pt in points_l:
        p_pt = tf_l.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.size = Pt(11)
        p_pt.font.color.rgb = C_TEXT
        p_pt.space_before = Pt(6)

    # Right Box: Universal Bypass & Patch
    c_r = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.7), Inches(5.0))
    c_r.fill.solid()
    c_r.fill.fore_color.rgb = C_CARD
    c_r.line.color.rgb = C_GREEN
    c_r.line.width = Pt(1.5)

    tf_r = c_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = Inches(0.25)

    p = tf_r.paragraphs[0]
    p.text = "Universal Bypass & Remediasi Kode"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    points_r = [
        "Payload Universal (Login Akun Pertama):\n' OR 1=1 LIMIT 1 #",
        "Remediasi Standar Industri (Prepared Statements):\n$stmt = $mysqli->prepare('SELECT id, fullname, role FROM users WHERE username = ? AND password = ?');\n$stmt->bind_param('ss', $username, $password);\n$stmt->execute();",
        "Prinsip Keamanan: Input diperlakukan murni sebagai parameter data, bukan instruksi logika SQL."
    ]
    for pt in points_r:
        p_pt = tf_r.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.size = Pt(11)
        p_pt.font.color.rgb = C_TEXT
        p_pt.space_before = Pt(6)

    # -----------------------------------------------------------------
    # SLIDE 4: SQLi Part 2 - UNION Data Exfiltration & PII Leak
    # -----------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "ATTACK VECTOR 2", "SQL Injection: UNION-Based Data Exfiltration", "Mengekstraksi NIK KTP, Gaji, Kartu Kredit, & Flag Rahasia dari MySQL")

    c_s4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.733), Inches(5.0))
    c_s4.fill.solid()
    c_s4.fill.fore_color.rgb = C_CARD
    c_s4.line.color.rgb = C_PURPLE
    c_s4.line.width = Pt(1.5)

    tf_s4 = c_s4.text_frame
    tf_s4.word_wrap = True
    tf_s4.margin_left = tf_s4.margin_right = tf_s4.margin_top = Inches(0.3)

    p = tf_s4.paragraphs[0]
    p.text = "Metodologi Ekstraksi Data Bocor (Real Data Breach)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE

    union_steps = [
        "1. Penentuan Jumlah Kolom & Tipe Data: Penyerang mencocokkan jumlah kolom tabel utama (5 kolom: id, fullname, username, email, role).",
        "2. Penyusunan Payload UNION SELECT:\n   ' UNION SELECT id, fullname, password, credit_card, api_secret_key FROM users #",
        "3. Ekstraksi Data Sensitif:\n   • Kartu Kredit Pimpinan: 4532-8921-3341-9012\n   • NIK KTP: 3271041982050001 | Gaji: Rp 45.000.000\n   • Secret Token Flag: FLAG{sql1_b0c0r_d4t4_kr3d3ns14l_m4st3r}",
        "4. Pembuktian Validitas: Mahasiswa dapat membuka tab Live DB Inspector atau phpMyAdmin untuk memverifikasi data yang bocor."
    ]
    for st in union_steps:
        p_st = tf_s4.add_paragraph()
        p_st.text = st
        p_st.font.size = Pt(12)
        p_st.font.color.rgb = C_TEXT
        p_st.space_before = Pt(8)

    # -----------------------------------------------------------------
    # SLIDE 5: XSS Triad (Reflected, Stored, DOM)
    # -----------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "ATTACK VECTOR 3", "Cross-Site Scripting (XSS) Triad", "Perbandingan 3 varian XSS pada SecLab")

    xss_cards = [
        ("Reflected XSS", "Non-Persistent", "Payload dikirim via parameter GET/POST (search bar) dan langsung dipantulkan di respon browser korban.", "<script>alert(document.domain)</script>", "htmlspecialchars($q, ENT_QUOTES, 'UTF-8')", C_YELLOW),
        ("Stored XSS", "Persistent (Tersimpan di MySQL)", "Payload masuk permanen ke tabel comments MySQL. Semua user yang melihat feed diskusi otomatis terinfeksi malware.", "<img src=x onerror=\"alert('Stored XSS Pwned')\">", "Sanitasi input + Contextual HTML Encoding", C_RED),
        ("DOM-Based XSS", "Client-Side Only", "Eksekusi terjadi murni di DOM engine browser (Source: location.hash -> Sink: innerHTML). Server MySQL tidak menerima traffic ini.", "#<img src=1 onerror=alert('DOM_XSS')>", "Ganti innerHTML dengan textContent", C_CYAN)
    ]

    for i, (name, tipe, desc, payload, patch, col) in enumerate(xss_cards):
        x = Inches(0.8 + i * 3.95)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), Inches(3.8), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = name
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col

        p_t = tf.add_paragraph()
        p_t.text = "[" + tipe + "]"
        p_t.font.size = Pt(11)
        p_t.font.color.rgb = C_MUTED
        p_t.space_before = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = C_TEXT
        p_d.space_before = Pt(6)

        p_pl = tf.add_paragraph()
        p_pl.text = "Payload Contoh:"
        p_pl.font.size = Pt(11)
        p_pl.font.bold = True
        p_pl.font.color.rgb = col
        p_pl.space_before = Pt(8)

        p_code = tf.add_paragraph()
        p_code.text = payload
        p_code.font.size = Pt(10)
        p_code.font.color.rgb = C_TEXT
        p_code.space_before = Pt(2)

        p_pch = tf.add_paragraph()
        p_pch.text = "Cara Patch:"
        p_pch.font.size = Pt(11)
        p_pch.font.bold = True
        p_pch.font.color.rgb = C_GREEN
        p_pch.space_before = Pt(8)

        p_pcode = tf.add_paragraph()
        p_pcode.text = patch
        p_pcode.font.size = Pt(10)
        p_pcode.font.color.rgb = C_TEXT
        p_pcode.space_before = Pt(2)

    # -----------------------------------------------------------------
    # SLIDE 6: Insecure File Upload to Web Shell RCE
    # -----------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "ATTACK VECTOR 4", "Insecure File Upload ke Remote Code Execution (RCE)", "Dari formulir upload avatar menuju pengambilalihan kendali server OS")

    c_s6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.733), Inches(5.0))
    c_s6.fill.solid()
    c_s6.fill.fore_color.rgb = C_CARD
    c_s6.line.color.rgb = C_RED
    c_s6.line.width = Pt(1.5)

    tf_s6 = c_s6.text_frame
    tf_s6.word_wrap = True
    tf_s6.margin_left = tf_s6.margin_right = tf_s6.margin_top = Inches(0.3)

    p = tf_s6.paragraphs[0]
    p.text = "Rantai Serangan Web Shell (Kill Chain)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_RED

    upload_flow = [
        "1. Pembuatan Web Shell (sample_materials/webshell.php):\n   <?php echo shell_exec($_GET['cmd']); ?>",
        "2. Upload Berkas Tanpa Whitelist Ekstensi: Aplikasi menyimpan berkas langsung ke direktori publik `/uploads/`.",
        "3. Trigger Eksekusi Perintah OS:\n   http://localhost/seclab-cybersecurity/uploads/webshell.php?cmd=whoami\n   http://localhost/seclab-cybersecurity/uploads/webshell.php?cmd=dir",
        "4. Defense-in-Depth (Pencegahan Berlapis):\n   • Whitelist Ekstensi Ketat: Hanya izinkan format [jpg, jpeg, png, pdf].\n   • Validasi MIME Magic Byte: Gunakan finfo_file() bukan sekadar header Content-Type klien.\n   • Randomize Nama Berkas: Ubah nama menjadi hash unik (bin2hex(random_bytes(16)) . '.' . $ext).\n   • Nonaktifkan script execution permission pada folder /uploads/."
    ]
    for uf in upload_flow:
        p_uf = tf_s6.add_paragraph()
        p_uf.text = uf
        p_uf.font.size = Pt(11)
        p_uf.font.color.rgb = C_TEXT
        p_uf.space_before = Pt(6)

    # -----------------------------------------------------------------
    # SLIDE 7: Dual-Mode Learning & Live DB Inspector
    # -----------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "LEARNING METHODOLOGY", "Dual-Mode Engine & Live Database Inspector", "Metodologi evaluasi ganda: Eksploitasi (Red Team) dan Remediasi (Blue Team)")

    boxes_s7 = [
        ("Vulnerable Mode (Red Team)", "Aplikasi mengeksekusi kueri langsung dan merender input tanpa sanitasi. Mahasiswa mempraktikkan eksploitasi dan pengumpulan flag.", C_RED),
        ("Secure Mode / Patched (Blue Team)", "Aplikasi menerapkan Prepared Statements, Whitelisting, dan HTML Output Encoding. Mahasiswa membuktikan eksploitasi gagal.", C_GREEN),
        ("Live DB Inspector (Auditing)", "Menampilkan isi tabel MySQL seclab_db secara langsung untuk memverifikasi integritas data dan payload tersimpan.", C_CYAN)
    ]

    for i, (title, desc, col) in enumerate(boxes_s7):
        x = Inches(0.8 + i * 3.95)
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), Inches(3.8), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT
        p2.space_before = Pt(10)

        p3 = tf.add_paragraph()
        p3.text = "Target Pembelajaran:"
        p3.font.size = Pt(13)
        p3.font.bold = True
        p3.font.color.rgb = col
        p3.space_before = Pt(14)

        if i == 0:
            sub_pts = ["Bypass login admin", "Dapatkan data kartu kredit", "Trigger XSS popup", "Upload web shell"]
        elif i == 1:
            sub_pts = ["Pahami prepared statement", "Terapkan output encoding", "Validasi file whitelist", "Uji ulang serangan"]
        else:
            sub_pts = ["Cek tabel users MySQL", "Audit tabel comments", "Sinkronisasi phpMyAdmin", "Validasi bukti laporan"]

        for sp in sub_pts:
            p_sp = tf.add_paragraph()
            p_sp.text = "✓ " + sp
            p_sp.font.size = Pt(11)
            p_sp.font.color.rgb = C_TEXT
            p_sp.space_before = Pt(4)

    # -----------------------------------------------------------------
    # SLIDE 8: Rundown Praktikum 150 Menit & Rubrik Penilaian
    # -----------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "PRACTICAL TIMELINE", "Rundown Praktikum Kelas (150 Menit)", "Alokasi waktu dan target capaian mahasiswa")

    timeline_s8 = [
        ("00 - 25 Menit", "Briefing Teori & Skenario Data Breach", "Review konsep OWASP Top 10, struktur tabel MySQL seclab_db, dan aturan lab.", C_CYAN),
        ("25 - 75 Menit", "Hands-on Exploitation (Red Team)", "Mahasiswa melakukan Auth Bypass login, ekstraksi kartu kredit/flag, XSS, dan Web Shell.", C_RED),
        ("75 - 115 Menit", "Analisis Kode & Secure Coding (Blue Team)", "Mahasiswa mempelajari implementasi patch di Secure Mode dan memverifikasi serangan gagal.", C_GREEN),
        ("115 - 140 Menit", "Penyusunan Laporan & Bukti phpMyAdmin", "Mengambil screenshot pembuktian dari Live DB Inspector dan phpMyAdmin.", C_PURPLE),
        ("140 - 150 Menit", "Evaluasi Dosen & Submission", "Review hasil temuan kelas dan penilaian capaian praktikum.", C_YELLOW)
    ]

    for i, (waktu, judul, sub, col) in enumerate(timeline_s8):
        y = Inches(1.9 + i * 1.0)
        box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(0.85))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD
        box.line.color.rgb = col
        box.line.width = Pt(1)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.12)
        tf.margin_left = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = f"{waktu}  |  {judul}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT
        p2.space_before = Pt(2)

    # Save to both target directories
    output_xampp = r"C:\xampp\htdocs\seclab-cybersecurity\Materi_Cybersecurity_SecLab.pptx"
    output_user = r"C:\Users\User\seclab-cybersecurity\Materi_Cybersecurity_SecLab.pptx"
    
    prs.save(output_xampp)
    prs.save(output_user)
    print("Updated presentation saved successfully at:")
    print(" -", output_xampp)
    print(" -", output_user)

if __name__ == "__main__":
    build_presentation()
